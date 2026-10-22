import tempfile
import unittest
from pathlib import Path

import openpyxl

from src.deduplication.diff_analyzer import DocumentDiffAnalyzer
from src.reporting.excel_exporter import ExcelExporter


class TestDocumentDiff(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_path = Path(self.temp_dir.name)

    def tearDown(self):
        try:
            self.temp_dir.cleanup()
        except Exception:
            pass

    def test_split_into_paragraphs(self):
        """Kiểm tra tách đoạn văn bản và loại trừ các nhãn kỹ thuật phân trang."""
        text = """--- [Trang 1: Bìa] ---
        CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM
        Độc lập - Tự do - Hạnh phúc

        BÁO CÁO KỸ THUẬT ĐỊA CHẤT

        --- [Trang 2: Nội dung] ---
        Phần 1: Tổng quan mỏ Bạch Hổ.
        Vị trí địa lý thuộc bể Cửu Long.
        """
        paras = DocumentDiffAnalyzer._split_into_paragraphs(text)
        self.assertNotIn("--- [Trang 1: Bìa] ---", paras)
        self.assertNotIn("--- [Trang 2: Nội dung] ---", paras)
        self.assertTrue(any("CỘNG HÒA" in p for p in paras))
        self.assertTrue(any("Phần 1: Tổng quan" in p for p in paras))

    def test_analyze_diff(self):
        """Kiểm tra phân tích diff trích xuất tóm tắt AI và đoạn thêm mới."""
        orig_text = (
            "TRƯỜNG ĐẠI HỌC BÁCH KHOA HÀ NỘI\n\n"
            "BÁO CÁO KẾT QUẢ THỰC TẬP TỐT NGHIỆP\n\n"
            "Chương 1: Giới thiệu chung về hệ thống khai thác dầu khí.\n\n"
            "Chương 2: Đánh giá lưu lượng dòng chảy tầng Mioxen.\n\n"
            "Kết luận: Hệ thống đạt hiệu suất cao."
        )

        dup_text = (
            "TRƯỜNG ĐẠI HỌC BÁCH KHOA HÀ NỘI\n\n"
            "BÁO CÁO KẾT QUẢ THỰC TẬP TỐT NGHIỆP - BẢN NHÁP ĐIỀU CHỈNH\n\n"
            "Chương 1: Giới thiệu chung về hệ thống khai thác dầu khí.\n\n"
            "Chương 2: Đánh giá lưu lượng dòng chảy tầng Mioxen và Oligoxen hạ.\n\n"
            "Bổ sung: Phân tích áp suất đáy giếng.\n\n"
            "Kết luận: Hệ thống đạt hiệu suất cao."
        )

        analyzer = DocumentDiffAnalyzer(llm_extractor=None)

        res = analyzer.analyze_diff(
            orig_text=orig_text,
            dup_text=dup_text,
            orig_name="Báo cáo gốc.pdf",
            dup_name="Báo cáo draft.docx",
            sim_pct=96.5,
        )

        self.assertIn("summary", res)
        self.assertEqual(res["similarity_pct"], 96.5)
        self.assertGreater(res["added_count"], 0)
        self.assertTrue(len(res["added_snippets"]) > 0)
        self.assertEqual(res["duplicate_file"], "Báo cáo draft.docx")
        self.assertEqual(res["original_file"], "Báo cáo gốc.pdf")

    def test_excel_diff_sheet_export(self):
        """Kiểm tra xuất Excel 2 sheet: Sheet 1 (Hover comment) và Sheet 2 (Đối chiếu Trùng lặp)."""
        data = [
            {
                "file_name": "Báo cáo chính thức.pdf",
                "file_path": str(self.base_path / "Báo cáo chính thức.pdf"),
                "is_duplicate": False,
                "metadata": {
                    "title": "Báo cáo Địa chất",
                    "contract_name": "Bạch Hổ",
                    "originator": "VPI",
                    "confidence_score": 98.0
                }
            },
            {
                "file_name": "Báo cáo chính thức-draft.docx",
                "file_path": str(self.base_path / "Báo cáo chính thức-draft.docx"),
                "is_duplicate": True,
                "duplicate_of": "Báo cáo chính thức.pdf",
                "metadata": {
                    "title": "Báo cáo Địa chất bản nháp",
                    "contract_name": "Bạch Hổ",
                    "originator": "VPI",
                    "confidence_score": 90.0,
                    "duplicate_of": "Báo cáo chính thức.pdf",
                    "diff_info": {
                        "duplicate_file": "Báo cáo chính thức-draft.docx",
                        "original_file": "Báo cáo chính thức.pdf",
                        "similarity_pct": 97.2,
                        "summary": "Bản nháp thêm 2 đoạn phân tích áp suất đáy giếng.",
                        "added_snippets": ["Phân tích áp suất đáy giếng BH-01"],
                        "deleted_snippets": [],
                        "html_path": str(self.base_path / "doi_chieu.html")
                    }
                }
            }
        ]

        excel_out = self.base_path / "tong_hop_tai_lieu_20261022.xlsx"
        ExcelExporter.export(data, excel_out)

        self.assertTrue(excel_out.exists())

        wb = openpyxl.load_workbook(excel_out)
        self.assertIn("Báo cáo Dầu khí", wb.sheetnames)
        self.assertIn("Đối chiếu Trùng lặp", wb.sheetnames)

        # 1. Kiểm tra Sheet 1: có Hover comment ở ô Trạng thái trùng lặp
        ws1 = wb["Báo cáo Dầu khí"]
        dup_cell = ws1.cell(row=3, column=16)  # Row 3 là Báo cáo chính thức-draft.docx, cột 16 là Trạng thái trùng lặp
        self.assertIn("Gần trùng lặp", str(dup_cell.value))
        self.assertIsNotNone(dup_cell.comment)
        self.assertIn("AI Diff", dup_cell.comment.text)
        self.assertIn("Bản nháp thêm 2 đoạn", dup_cell.comment.text)

        # 2. Kiểm tra Sheet 2: Bảng Đối chiếu Trùng lặp
        ws2 = wb["Đối chiếu Trùng lặp"]
        headers = [ws2.cell(row=1, column=c).value for c in range(1, ws2.max_column + 1)]
        expected_headers = [
            "STT",
            "Tên tệp trùng lặp",
            "Tệp gốc đối chiếu",
            "Độ tương đồng",
            "Tóm tắt sai khác (AI Diff Summary)",
            "Trích đoạn Thêm mới (+)",
            "Trích đoạn Bị cắt bỏ (-)",
        ]
        self.assertEqual(headers, expected_headers)

        self.assertEqual(ws2.cell(row=2, column=2).value, "Báo cáo chính thức-draft.docx")
        self.assertEqual(ws2.cell(row=2, column=3).value, "Báo cáo chính thức.pdf")
        self.assertEqual(ws2.cell(row=2, column=4).value, "97.2%")
        self.assertIn("áp suất đáy giếng", ws2.cell(row=2, column=5).value)

        wb.close()


if __name__ == "__main__":
    unittest.main()
