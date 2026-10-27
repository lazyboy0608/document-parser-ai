import shutil
import tempfile
import time
import unittest
from pathlib import Path
import openpyxl

from src.ingestion.folder_watcher import HotFolderWatcher
from src.utils.state_tracker import StateTracker
from src.reporting.excel_exporter import ExcelExporter


class TestAutoExcelSyncWatcher(unittest.TestCase):
    """
    Kiểm thử tính năng Zero-Click Auto-Sync:
    Khi HotFolderWatcher đang chạy, người dùng sửa file Excel và bấm Lưu (Ctrl+S),
    Watcher tự động bắt sự kiện, đồng bộ ngược vào SQLite và làm mới file Excel.
    """

    def setUp(self):
        self.test_dir = Path(tempfile.mkdtemp())
        self.watch_dir = self.test_dir / "hot_folder"
        self.output_dir = self.test_dir / "output"
        self.db_path = self.test_dir / "test_pipeline.db"
        self.watch_dir.mkdir(parents=True, exist_ok=True)
        self.output_dir.mkdir(parents=True, exist_ok=True)

        self.watcher = HotFolderWatcher(
            watch_dir=self.watch_dir,
            output_dir=self.output_dir,
            stable_seconds=0.5,
            mode="copy",
            db_path=self.db_path,
        )

    def tearDown(self):
        if getattr(self, "watcher", None):
            self.watcher.stop()
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_auto_sync_on_excel_save(self):
        """Mô phỏng thao tác sửa file Excel và lưu (Ctrl+S) -> Kiểm tra CSDL tự động cập nhật."""
        date_str = "20261027"
        daily_output = self.output_dir / date_str
        daily_output.mkdir(parents=True, exist_ok=True)

        # 1. Tạo một tệp tài liệu giả lập trong thư mục output
        doc_dir = daily_output / "Contract_Old"
        doc_dir.mkdir(parents=True, exist_ok=True)
        sample_file = doc_dir / "bao_cao_dia_chat.pdf"
        sample_file.write_text("Noi dung dia chat dau khi", encoding="utf-8")

        # Lưu bản ghi ban đầu vào SQLite
        initial_meta = {
            "title": "Báo cáo Địa chất",
            "contract_name": "Lô Cũ",
            "author": "Kỹ sư A",
            "confidence_score": 75.0,
            "low_confidence_fields": ["contract_name"],
            "field_warnings": {"contract_name": "Chữ mờ"}
        }
        self.watcher.pipeline.state_tracker.save_state(
            file_path=str(sample_file),
            file_hash="hash_123456",
            status="COMPLETED",
            is_duplicate=False,
            metadata=initial_meta
        )

        # 2. Xuất file Excel ban đầu
        excel_path = daily_output / f"tong_hop_tai_lieu_{date_str}.xlsx"
        ExcelExporter.export([
            {
                "file_name": sample_file.name,
                "file_path": str(sample_file),
                "is_duplicate": False,
                "metadata": initial_meta
            }
        ], excel_path)
        self.assertTrue(excel_path.exists())

        # 3. Khởi động Watcher (không chạy scan ban đầu của watch_dir để tập trung kiểm tra output_dir)
        # Bắt đầu observer và worker loops
        event_handler = self.watcher.observer._emitter_for_path[str(self.watch_dir)] if hasattr(self.watcher, "_emitter_for_path") else None
        
        # Gọi start của watcher trong luồng riêng
        import threading
        watcher_thread = threading.Thread(target=lambda: self.watcher.start(run_initial_scan=False), daemon=True)
        watcher_thread.start()
        time.sleep(1.0) # Đợi observer sẵn sàng

        # 4. Người dùng mở file Excel lên sửa: 'Lô Cũ' -> 'Lô 09-1 (Bạch Hổ)' và bấm Lưu (Ctrl + S)
        wb = openpyxl.load_workbook(excel_path)
        ws = wb["Báo cáo Dầu khí"]
        # Cột 4 là CONTRACT NAME (STT=1, Tên tệp=2, Well link=3, CONTRACT NAME=4)
        ws.cell(row=2, column=4).value = "Lô 09-1 (Bạch Hổ)"
        # Sửa thêm Tác giả ở Cột 9 (AUTHOR)
        ws.cell(row=2, column=9).value = "Nguyễn Văn Chuyên Gia"
        wb.save(excel_path)
        wb.close()

        # 5. Đợi Watcher tự động bắt sự kiện và đồng bộ ngầm (chờ tối đa 10 giây)
        found = False
        matching_rec = None
        for _ in range(20):
            time.sleep(0.5)
            all_completed = self.watcher.pipeline.state_tracker.get_all_completed()
            for fp, rec in all_completed.items():
                if "bao_cao_dia_chat.pdf" in fp:
                    if rec.get("metadata", {}).get("contract_name") == "Lô 09-1 (Bạch Hổ)":
                        matching_rec = rec
                        found = True
                        break
            if found:
                break

        self.assertTrue(found, "Watcher không hoàn tất tự động đồng bộ trong thời gian chờ!")
        self.assertIsNotNone(matching_rec)
        meta_updated = matching_rec.get("metadata", {})

        # Xác nhận giá trị đã được cập nhật chuẩn xác
        self.assertEqual(meta_updated.get("contract_name"), "Lô 09-1 (Bạch Hổ)")
        self.assertEqual(meta_updated.get("author"), "Nguyễn Văn Chuyên Gia")

        # Độ tin cậy của trường con người sửa phải đạt 100%
        self.assertEqual(meta_updated.get("field_confidences", {}).get("contract_name"), 100.0)
        self.assertNotIn("contract_name", meta_updated.get("low_confidence_fields", []))

        # Kiểm tra file đã được tự động tái phân cụm (Re-clustered) sang thư mục mới
        new_path = Path(matching_rec["file_path"])
        self.assertTrue(new_path.exists(), f"File mới không tồn tại: {new_path}")
        self.assertIn("Lô_09-1", new_path.parent.name)

        # Dừng watcher
        self.watcher.stop()


if __name__ == "__main__":
    unittest.main()
