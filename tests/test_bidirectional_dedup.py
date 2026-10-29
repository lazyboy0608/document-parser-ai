import shutil
import tempfile
import unittest
from pathlib import Path
import openpyxl

from src.utils.state_tracker import StateTracker
from src.clustering.file_organizer import FileOrganizer
from src.deduplication.pipeline_dedup import DeduplicationPipeline
from src.reporting.excel_exporter import ExcelExporter


class TestBidirectionalDeduplication(unittest.TestCase):
    """
    Kiểm thử cơ chế liên kết trùng lặp hai chiều (Bi-directional Duplicate Grouping):
    - Tệp nạp đầu tiên A được hệ thống nhận diện là gốc ban đầu.
    - Tệp nạp sau B trùng hoặc gần trùng với A.
    - Khi B nạp vào:
      1. B được gắn cờ is_duplicate=True, duplicate_of="A.docx", có chung group_id với A.
      2. Tệp đầu tiên A phải được CẬP NHẬT NGƯỢC (Reverse Link):
         - duplicate_variants chứa ["B.docx"]
         - duplicate_status hiển thị rõ "Có bản trùng lặp (với B.docx)"
         - Có chung group_id với B.
      3. Trên tệp Excel xuất ra:
         - Ô trạng thái của tệp A hiển thị "Có bản trùng lặp (với B.docx)"
         - Hover Comment của tệp A cảnh báo về các tệp phái sinh nạp sau.
         - Sheet 2 "Đối chiếu Trùng lặp" chứa mã Group ID.
    """

    def setUp(self):
        self.test_dir = Path(tempfile.mkdtemp())
        self.output_dir = self.test_dir / "output"
        self.output_dir.mkdir(parents=True, exist_ok=True)

        self.db_path = self.test_dir / "test_pipeline.db"
        self.state_tracker = StateTracker(self.db_path)
        self.file_organizer = FileOrganizer(self.output_dir)

        self.dedup_pipeline = DeduplicationPipeline(
            state_tracker=self.state_tracker,
            file_organizer=self.file_organizer,
            mode="copy",  # copy để giữ lại file nguồn test
        )

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_reverse_link_on_exact_duplicate(self):
        """Kiểm tra cập nhật ngược tệp đầu khi có tệp trùng 100% SHA-256 nạp sau."""
        # 1. Giả lập tệp A nạp đầu tiên
        file_a = self.test_dir / "BaoCaoGoc.docx"
        file_a.write_text("Noi dung ky thuat dau khi giong nhau hoan toan 100%", encoding="utf-8")

        # Đăng ký tệp A là tệp gốc đã xử lý trước đó
        meta_a = {
            "title": "Báo cáo Kỹ thuật Mỏ",
            "contract_name": "Lô 09-1",
            "confidence_score": 95.0,
            "duplicate_status": "Gốc (Không trùng)"
        }
        hash_a = self.dedup_pipeline.exact_dedup.calculate_file_hash(file_a)
        self.state_tracker.save_state(
            file_path=str(file_a),
            file_hash=hash_a,
            status="COMPLETED",
            is_duplicate=False,
            metadata=meta_a
        )
        self.dedup_pipeline.exact_dedup.register_file(file_a, hash_a)

        # 2. Tệp B nạp sau, nội dung y hệt A (exact duplicate)
        file_b = self.test_dir / "BaoCaoGoc_Copy.docx"
        file_b.write_text("Noi dung ky thuat dau khi giong nhau hoan toan 100%", encoding="utf-8")

        res_b = self.dedup_pipeline.process(file_b, date_str="20261023")

        # Kiểm tra tệp B
        self.assertTrue(res_b.is_duplicate)
        self.assertEqual(res_b.matched_orig_name, "BaoCaoGoc.docx")
        self.assertIsNotNone(res_b.group_id)

        # 3. KIỂM TRA TỆP ĐẦU TIÊN (A) ĐÃ ĐƯỢC CẬP NHẬT NGƯỢC (REVERSE LINK) TRONG SQLITE
        completed_records = self.state_tracker.get_all_completed()
        norm_a = self.state_tracker._norm(str(file_a))
        rec_a = completed_records.get(norm_a)

        self.assertIsNotNone(rec_a)
        meta_a_updated = rec_a.get("metadata", {})
        self.assertTrue(meta_a_updated.get("has_duplicates"))
        self.assertIn("BaoCaoGoc_Copy.docx", meta_a_updated.get("duplicate_variants", []))
        self.assertIn("Có bản trùng lặp (với BaoCaoGoc_Copy.docx)", meta_a_updated.get("duplicate_status"))
        self.assertEqual(meta_a_updated.get("duplicate_group_id"), res_b.group_id)

    def test_reverse_link_on_semantic_duplicate(self):
        """Kiểm tra cập nhật ngược tệp đầu khi có tệp phái sinh nạp sau (SimHash >= 95%)."""
        # Tạo văn bản gốc
        base_text = "Tập đoàn Dầu khí Việt Nam báo cáo tiến độ khai thác giếng khoan BK-20 mỏ Bạch Hổ. " * 30
        file_orig = self.test_dir / "BaoCaoKhaiThac.txt"
        file_orig.write_text(base_text, encoding="utf-8")

        meta_orig = {
            "title": "Báo cáo Khai thác Mỏ Bạch Hổ",
            "contract_name": "Lô 09-1",
            "confidence_score": 92.0,
            "duplicate_status": "Gốc (Không trùng)"
        }
        shash_orig = self.dedup_pipeline.simhash_dedup.compute_simhash(base_text)
        fhash_orig = self.dedup_pipeline.exact_dedup.calculate_file_hash(file_orig)

        self.state_tracker.save_state(
            file_path=str(file_orig),
            file_hash=fhash_orig,
            status="COMPLETED",
            is_duplicate=False,
            metadata=meta_orig,
            simhash=shash_orig
        )
        self.dedup_pipeline.active_simhashes[str(file_orig)] = shash_orig
        self.dedup_pipeline.exact_dedup.register_file(file_orig, fhash_orig)

        # Tạo bản draft nạp sau (chỉ thêm 1 câu nhỏ, độ tương đồng >= 97%)
        draft_text = base_text + "\nBổ sung thông số áp suất đáy giếng đạt 250 atm ngày 23/10."
        file_draft = self.test_dir / "BaoCaoKhaiThac_draft.txt"
        file_draft.write_text(draft_text, encoding="utf-8")

        res_draft = self.dedup_pipeline.process(file_draft, date_str="20261023")

        self.assertTrue(res_draft.is_duplicate)
        self.assertEqual(res_draft.matched_orig_name, "BaoCaoKhaiThac.txt")
        self.assertIsNotNone(res_draft.group_id)

        # Kiểm tra tệp gốc được cập nhật ngược
        completed = self.state_tracker.get_all_completed()
        norm_orig = self.state_tracker._norm(str(file_orig))
        rec_orig = completed.get(norm_orig)

        self.assertIsNotNone(rec_orig)
        meta_updated = rec_orig.get("metadata", {})
        self.assertIn("BaoCaoKhaiThac_draft.txt", meta_updated.get("duplicate_variants", []))
        self.assertIn("Có bản trùng lặp", meta_updated.get("duplicate_status"))
        self.assertEqual(meta_updated.get("duplicate_group_id"), res_draft.group_id)

        # 4. Kiểm tra xuất Excel
        excel_path = self.output_dir / "test_report.xlsx"
        items_to_export = [
            rec_orig,
            {
                "file_name": file_draft.name,
                "file_path": str(file_draft),
                "is_duplicate": True,
                "duplicate_of": "BaoCaoKhaiThac.txt",
                "duplicate_group_id": res_draft.group_id,
                "metadata": res_draft.inherited_metadata,
            }
        ]
        ExcelExporter.export(items_to_export, excel_path)
        self.assertTrue(excel_path.exists())

        # Mở kiểm tra cell comment và nội dung
        wb = openpyxl.load_workbook(excel_path)
        ws1 = wb["Báo cáo Dầu khí"]

        # Tìm ô Trạng thái trùng lặp của tệp gốc (hàng 2)
        dup_col_idx = ExcelExporter.COLUMNS.index("Trạng thái trùng lặp") + 1
        cell_status_orig = ws1.cell(row=2, column=dup_col_idx)
        self.assertIn("Có bản trùng lặp (với BaoCaoKhaiThac_draft.txt)", str(cell_status_orig.value))
        self.assertIsNotNone(cell_status_orig.comment)
        self.assertIn("CẢNH BÁO TÀI LIỆU GỐC CÓ BẢN PHÁI SINH", cell_status_orig.comment.text)
        self.assertIn("BaoCaoKhaiThac_draft.txt", cell_status_orig.comment.text)

        # Kiểm tra Sheet 2 "Đối chiếu Trùng lặp" có cột Group ID
        ws2 = wb["Đối chiếu Trùng lặp"]
        col_group_id = ws2.cell(row=1, column=2).value
        self.assertEqual(col_group_id, "Mã nhóm (Group ID)")
        val_group_id = ws2.cell(row=2, column=2).value
        self.assertEqual(val_group_id, res_draft.group_id)

        wb.close()


if __name__ == "__main__":
    unittest.main()
