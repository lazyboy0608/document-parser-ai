import json
import sqlite3
import tempfile
import unittest
from pathlib import Path

import openpyxl
import pandas as pd

from src.reporting.excel_exporter import ExcelExporter
from src.reporting.sync_engine import ExcelSyncEngine
from src.utils.state_tracker import StateTracker


class TestExcelReverseSync(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_path = Path(self.temp_dir.name)
        self.db_path = self.base_path / "test_pipeline.db"
        self.output_dir = self.base_path / "output"
        self.output_dir.mkdir(parents=True, exist_ok=True)

        self.state_tracker = StateTracker(self.db_path)
        self.sync_engine = ExcelSyncEngine(db_path=self.db_path, output_dir=self.output_dir)

    def tearDown(self):
        try:
            self.temp_dir.cleanup()
        except Exception:
            pass

    def test_sync_orange_low_confidence_cell(self):
        """Kiểm tra sửa ô màu cam (độ tin cậy thấp) -> Cập nhật 100% tin cậy, gỡ highlight & warning."""
        date_str = "20261022"
        cluster_dir = self.output_dir / date_str / "Bạch_Hổ"
        cluster_dir.mkdir(parents=True, exist_ok=True)
        file_path = cluster_dir / "report_01.pdf"
        file_path.write_bytes(b"Dummy PDF content")

        initial_meta = {
            "title": "Chưa rõ",
            "author": "Nguyễn Văn A",
            "originator": "VPI",
            "contractor": "PetroVietnam",
            "contract_name": "Bạch Hổ",
            "well_link": "BH-01",
            "public_date": "2026-01-01",
            "language": "Vietnamese",
            "document_type": "Report",
            "encl": "No",
            "encl_name": "None",
            "confidence_score": 65.0,
            "field_confidences": {"title": 40.0, "author": 90.0, "contract_name": 90.0},
            "low_confidence_fields": ["title"],
            "field_warnings": {"title": "Chữ mờ nhòe do chất lượng scan thấp"},
            "verified_fields": []
        }

        self.state_tracker.save_state(
            str(file_path),
            file_hash="hash123",
            status="COMPLETED",
            is_duplicate=False,
            metadata=initial_meta
        )

        excel_path = self.output_dir / date_str / f"tong_hop_tai_lieu_{date_str}.xlsx"
        ExcelExporter.export([{
            "file_name": file_path.name,
            "file_path": str(file_path),
            "is_duplicate": False,
            "metadata": initial_meta
        }], excel_path)

        # 1. Kiểm tra ban đầu ô Title có comment cảnh báo
        wb = openpyxl.load_workbook(excel_path)
        ws = wb["Báo cáo Dầu khí"]
        # Cột TITLE là cột 5
        title_cell = ws.cell(row=2, column=5)
        self.assertEqual(title_cell.value, "Chưa rõ")
        self.assertIsNotNone(title_cell.comment)
        self.assertIn("Chữ mờ nhòe", title_cell.comment.text)
        wb.close()

        # 2. Con người sửa ô TITLE trên Excel thành "Báo cáo Khai thác Mỏ Bạch Hổ"
        wb = openpyxl.load_workbook(excel_path)
        ws = wb["Báo cáo Dầu khí"]
        ws.cell(row=2, column=5, value="Báo cáo Khai thác Mỏ Bạch Hổ")
        wb.save(excel_path)
        wb.close()

        # 3. Đồng bộ ngược từ Excel về DB
        sync_res = self.sync_engine.sync_from_excel(excel_path)
        self.assertEqual(sync_res["updated_files_count"], 1)

        # 4. Kiểm tra CSDL SQLite đã cập nhật
        with sqlite3.connect(self.db_path) as conn:
            cur = conn.cursor()
            cur.execute("SELECT metadata_json FROM processed_files WHERE file_path = ?", (str(file_path),))
            row = cur.fetchone()
            self.assertIsNotNone(row)
            updated_meta = json.loads(row[0])

            self.assertEqual(updated_meta["title"], "Báo cáo Khai thác Mỏ Bạch Hổ")
            self.assertEqual(updated_meta["field_confidences"]["title"], 100.0)
            self.assertNotIn("title", updated_meta["low_confidence_fields"])
            self.assertNotIn("title", updated_meta["field_warnings"])
            self.assertIn("title", updated_meta["verified_fields"])
            self.assertGreater(updated_meta["confidence_score"], 65.0)

    def test_sync_normal_cell_override(self):
        """Kiểm tra con người sửa ô bình thường (không tô cam) -> Vẫn được công nhận Ground Truth 100%."""
        date_str = "20261022"
        cluster_dir = self.output_dir / date_str / "General"
        cluster_dir.mkdir(parents=True, exist_ok=True)
        file_path = cluster_dir / "doc_normal.pdf"
        file_path.write_bytes(b"Dummy normal document")

        initial_meta = {
            "title": "Báo cáo Tài chính",
            "author": "Kế toán viên A",
            "originator": "PVN",
            "contractor": "PVN",
            "contract_name": "Chung",
            "well_link": "N/A",
            "public_date": "2026-03-15",
            "language": "Vietnamese",
            "document_type": "Report",
            "encl": "No",
            "encl_name": "None",
            "confidence_score": 95.0,
            "field_confidences": {"author": 90.0, "title": 95.0},
            "low_confidence_fields": [],
            "field_warnings": {},
            "verified_fields": []
        }

        self.state_tracker.save_state(
            str(file_path),
            file_hash="hash456",
            status="COMPLETED",
            is_duplicate=False,
            metadata=initial_meta
        )

        excel_path = self.output_dir / date_str / f"tong_hop_tai_lieu_{date_str}.xlsx"
        ExcelExporter.export([{
            "file_name": file_path.name,
            "file_path": str(file_path),
            "is_duplicate": False,
            "metadata": initial_meta
        }], excel_path)

        # Con người sửa Tác giả từ "Kế toán viên A" -> "Nguyễn Thị Mai (Trưởng phòng)"
        wb = openpyxl.load_workbook(excel_path)
        ws = wb["Báo cáo Dầu khí"]
        # Cột AUTHOR là cột 9
        ws.cell(row=2, column=9, value="Nguyễn Thị Mai (Trưởng phòng)")
        wb.save(excel_path)
        wb.close()

        sync_res = self.sync_engine.sync_from_excel(excel_path)
        self.assertEqual(sync_res["updated_files_count"], 1)

        with sqlite3.connect(self.db_path) as conn:
            cur = conn.cursor()
            cur.execute("SELECT metadata_json FROM processed_files WHERE file_path = ?", (str(file_path),))
            updated_meta = json.loads(cur.fetchone()[0])
            self.assertEqual(updated_meta["author"], "Nguyễn Thị Mai (Trưởng phòng)")
            self.assertEqual(updated_meta["field_confidences"]["author"], 100.0)
            self.assertIn("author", updated_meta["verified_fields"])

    def test_sync_auto_recluster(self):
        """Kiểm tra con người sửa thông tin Dự án / Giếng khoan -> Tự động di chuyển thư mục phân cụm."""
        date_str = "20261022"
        old_cluster_dir = self.output_dir / date_str / "Viện_Dầu_khí"
        old_cluster_dir.mkdir(parents=True, exist_ok=True)
        file_path = old_cluster_dir / "nghiencuu.docx"
        file_path.write_bytes(b"Docx content")

        initial_meta = {
            "title": "Nghiên cứu địa chất",
            "author": "Nhóm nghiên cứu",
            "originator": "Viện Dầu khí",
            "contractor": "N/A",
            "contract_name": "N/A",
            "well_link": "N/A",
            "public_date": "2026-05-10",
            "language": "Vietnamese",
            "document_type": "Report",
            "encl": "No",
            "encl_name": "None",
            "confidence_score": 85.0,
            "field_confidences": {},
            "low_confidence_fields": [],
            "field_warnings": {},
            "verified_fields": []
        }

        self.state_tracker.save_state(
            str(file_path),
            file_hash="hash789",
            status="COMPLETED",
            is_duplicate=False,
            metadata=initial_meta
        )

        excel_path = self.output_dir / date_str / f"tong_hop_tai_lieu_{date_str}.xlsx"
        ExcelExporter.export([{
            "file_name": file_path.name,
            "file_path": str(file_path),
            "is_duplicate": False,
            "metadata": initial_meta
        }], excel_path)

        # Sửa CONTRACT NAME thành "Lô 05-1a"
        wb = openpyxl.load_workbook(excel_path)
        ws = wb["Báo cáo Dầu khí"]
        # CONTRACT NAME là cột 4
        ws.cell(row=2, column=4, value="Lô 05-1a")
        wb.save(excel_path)
        wb.close()

        sync_res = self.sync_engine.sync_from_excel(excel_path, auto_recluster=True)
        self.assertEqual(sync_res["updated_files_count"], 1)
        self.assertEqual(len(sync_res["reclustered_files"]), 1)

        # Kiểm tra file đã được di chuyển sang thư mục mới
        expected_new_path = self.output_dir / date_str / "Lô_05-1a" / "nghiencuu.docx"
        self.assertTrue(expected_new_path.exists())
        self.assertFalse(file_path.exists())

        # Kiểm tra DB lưu đúng đường dẫn mới
        with sqlite3.connect(self.db_path) as conn:
            cur = conn.cursor()
            cur.execute("SELECT file_path, metadata_json FROM processed_files WHERE file_path = ?", (str(expected_new_path),))
            row = cur.fetchone()
            self.assertIsNotNone(row)
            meta = json.loads(row[1])
            self.assertEqual(meta["contract_name"], "Lô 05-1a")


if __name__ == "__main__":
    unittest.main()
